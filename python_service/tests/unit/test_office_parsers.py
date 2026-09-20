"""Unit tests for the Office document parsers (Word support + PPT fallback)."""

from __future__ import annotations

import shutil
import subprocess

import pytest

from httpserver.services.office_service import OfficeService


@pytest.fixture(scope="module")
def service() -> OfficeService:
    return OfficeService()


def _make_docx(path, paragraphs, table=None):
    from docx import Document

    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        table_doc = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                table_doc.rows[r].cells[c].text = value
    document.save(str(path))


def test_parse_docx_paragraphs_and_table(service, tmp_path):
    docx_path = tmp_path / "sample.docx"
    _make_docx(
        docx_path,
        paragraphs=["第一段", "第二段"],
        table=[["姓名", "年龄"], ["张三", "30"]],
    )

    content = service._parse_docx(str(docx_path))
    assert "第一段" in content
    assert "第二段" in content
    assert "| 姓名 | 年龄 |" in content
    assert "| 张三 | 30 |" in content


def test_parse_docx_empty_document(service, tmp_path):
    docx_path = tmp_path / "empty.docx"
    _make_docx(docx_path, paragraphs=[])
    assert service._parse_docx(str(docx_path)) == ""


def test_parse_docx_corrupt_file_returns_error_string(service, tmp_path):
    bad = tmp_path / "bad.docx"
    bad.write_bytes(b"not a zip")
    content = service._parse_docx(str(bad))
    assert content.startswith("Error parsing DOCX file:")


@pytest.mark.skipif(shutil.which("antiword") is None, reason="antiword not installed")
def test_parse_doc_missing_file_reports_error(service, tmp_path):
    content = service._parse_doc(str(tmp_path / "missing.doc"))
    assert content.startswith("Error parsing DOC:")


@pytest.mark.asyncio
async def test_parse_file_text_content_with_doc_extension(service, tmp_path):
    """Deleted-file carving often yields text in a .doc record: preview it."""
    fake = tmp_path / "carved.doc"
    fake.write_text("Const ERROR_SUCCESS=0\r\nConst ERROR_CANNOT_COPY=266\r\n", encoding="ascii")
    content = await service.parse_file(str(fake))
    assert content.startswith("[注意]")
    assert "ERROR_SUCCESS" in content


@pytest.mark.asyncio
async def test_parse_file_partial_carved_text_previews_prefix(service, tmp_path):
    """A readable fragment followed by overwritten garbage shows the fragment."""
    import os

    fake = tmp_path / "carved.doc"
    raw = b"Const ERROR_SUCCESS=0\r\n" * 40 + os.urandom(4096)
    fake.write_bytes(raw)
    content = await service.parse_file(str(fake))
    assert "仅前" in content
    assert "ERROR_SUCCESS" in content


@pytest.mark.asyncio
async def test_parse_file_binary_mismatch_reports_unpreviewable(service, tmp_path):
    fake = tmp_path / "carved.docx"
    fake.write_bytes(b"\x00\x01\x02\x03" * 2000)
    content = await service.parse_file(str(fake))
    assert "无法提取可读文本" in content


@pytest.mark.asyncio
async def test_parse_file_real_docx_still_dispatches(service, tmp_path):
    docx_path = tmp_path / "real.docx"
    _make_docx(docx_path, paragraphs=["真实文档内容"])
    content = await service.parse_file(str(docx_path))
    assert "真实文档内容" in content


def test_parse_ppt_ole_extracts_text(service, monkeypatch):
    """The pure-Python PPT text walk parses TextChars/TextBytes atoms."""
    import struct
    import types

    # Build a minimal "PowerPoint Document" stream: one container record
    # wrapping a TextCharsAtom (UTF-16) and a TextBytesAtom (GBK).
    def atom(rec_type: int, payload: bytes) -> bytes:
        return struct.pack("<HHI", 0, rec_type, len(payload)) + payload

    children = atom(0x0FA0, "幻灯片标题".encode("utf-16-le")) + atom(
        0x0FA8, "正文内容".encode("gbk")
    )
    stream = struct.pack("<HHI", 0xF, 0x1000, len(children)) + children

    class FakeStream:
        def read(self):
            return stream

    class FakeOle:
        def exists(self, name):
            return True

        def openstream(self, name):
            return FakeStream()

    class FakeOleFileIO:
        def __init__(self, path):
            pass

        def __enter__(self):
            return FakeOle()

        def __exit__(self, *args):
            return None

    fake_module = types.SimpleNamespace(OleFileIO=FakeOleFileIO)
    monkeypatch.setitem(__import__("sys").modules, "olefile", fake_module)

    content = service._parse_ppt_ole("fake.ppt")
    assert "幻灯片标题" in content
    assert "正文内容" in content
