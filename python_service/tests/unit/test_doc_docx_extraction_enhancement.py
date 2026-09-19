"""Unit tests for the enhanced legacy-Office (.doc/.docx) extraction path.

Covers the screenshot-document failure mode: markitdown cannot convert legacy
.doc at all and antiword then emits only [pic] placeholders, so the pipeline
feeds placeholder noise to the LLM as evidence. The extractors now carve the
embedded images and OCR them (RapidOCR), and the DOCX walker covers tables,
headers/footers, text boxes and encrypted-compound diagnostics.
"""

import asyncio
import io
import struct
import zipfile

import pytest

from httpserver.services.extractors import doc_ocr
from httpserver.services.extractors.doc_ocr import (
    EmbeddedImage,
    build_image_ocr_markdown,
    carve_images_from_stream,
    has_meaningful_text,
)
from httpserver.services.extractors.office import DocExtractorProxy, DocxExtractor


def _png_blob(width=100, height=80, color=(200, 30, 30)) -> bytes:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (width, height), color).save(buf, format="PNG")
    return buf.getvalue()


def _jpeg_blob(width=120, height=90) -> bytes:
    from PIL import Image

    buf = io.BytesIO()
    Image.new("RGB", (width, height), (10, 90, 200)).save(buf, format="JPEG")
    return buf.getvalue()


def _image(width=100, height=80) -> EmbeddedImage:
    blob = _png_blob(width, height)
    import hashlib

    return EmbeddedImage(blob=blob, width=width, height=height, sha256=hashlib.sha256(blob).hexdigest())


# ---------------------------------------------------------------------------
# Placeholder gate
# ---------------------------------------------------------------------------

class TestMeaningfulTextGate:
    def test_pure_pic_lines_is_not_meaningful(self):
        assert not has_meaningful_text("[pic]\n[pic]\n[pic]\n")

    def test_repeated_pics_on_one_line_is_not_meaningful(self):
        assert not has_meaningful_text("[pic][pic]\n[pic]\n")

    def test_real_text_is_meaningful(self):
        assert has_meaningful_text("第二天\n性格灌输，家庭摸底\n讲故事灌输人物性格")

    def test_short_caption_alone_is_not_meaningful(self):
        assert not has_meaningful_text("[pic]\n充值记录\n")

    def test_empty(self):
        assert not has_meaningful_text("")
        assert not has_meaningful_text(None)


# ---------------------------------------------------------------------------
# Image carving
# ---------------------------------------------------------------------------

class TestCarveImages:
    def test_carves_png_and_jpeg_with_dedupe_and_size_filter(self):
        png = _png_blob()
        small_png = _png_blob(24, 24)  # icon-sized: filtered out
        jpg = _jpeg_blob()
        junk = b"\x00" * 64
        stream = junk + png + junk + png + junk + small_png + jpg + junk

        images = carve_images_from_stream(stream)
        kinds = []
        for image in images:
            kinds.append(image.blob[:8])

        assert len(images) == 2  # duplicate PNG dropped, tiny icon dropped
        assert images[0].blob == png
        assert images[1].blob == jpg
        assert images[0].width == 100 and images[0].height == 80
        assert len({image.sha256 for image in images}) == 2

    def test_empty_stream(self):
        assert carve_images_from_stream(b"") == []

    def test_truncated_png_is_dropped(self):
        png = _png_blob()
        images = carve_images_from_stream(png[:-20])  # no IEND -> not carved
        assert images == []


# ---------------------------------------------------------------------------
# OCR markdown rendering (engine stubbed — no model in unit tests)
# ---------------------------------------------------------------------------

class TestImageOcrMarkdown:
    def test_engine_unavailable_yields_diagnostic(self, monkeypatch):
        monkeypatch.setattr(doc_ocr, "get_ocr_engine", lambda: None)
        md = build_image_ocr_markdown([_image(), _image()], "x.doc")
        assert "2 张内嵌图片" in md
        assert "OCR" in md

    def test_no_images_yields_no_content_marker(self):
        assert "[No content extracted from" in build_image_ocr_markdown([], "x.doc")

    def test_sections_rendered_per_image(self, monkeypatch):
        monkeypatch.setattr(doc_ocr, "get_ocr_engine", lambda: object())
        monkeypatch.setattr(
            doc_ocr,
            "ocr_image_bytes",
            lambda blob: ["晚上7:02", "嗨，Ales，我叫李婷"],
        )
        md = build_image_ocr_markdown([_image(), _image(111, 91)], "x.doc")
        assert "## 图片 1（100x80）" in md
        assert "## 图片 2（111x91）" in md
        assert "我叫李婷" in md
        assert "[pic]" not in md


# ---------------------------------------------------------------------------
# DOC proxy: antiword text layer + image fallback
# ---------------------------------------------------------------------------

class TestDocExtractorProxy:
    def test_text_document_short_circuits_without_image_work(self, monkeypatch):
        calls = {}
        monkeypatch.setattr(
            DocExtractorProxy, "_run_antiword", staticmethod(lambda p: ("充值成功记录文本 " * 10, False))
        )

        def _fail(p):
            calls["carve"] = True
            raise AssertionError("image carving must not run for text documents")

        monkeypatch.setattr(doc_ocr, "extract_doc_images", _fail)
        md = asyncio.run(DocExtractorProxy().extract_to_markdown("fake.doc"))
        assert "充值成功记录文本" in md
        assert "carve" not in calls

    def test_placeholder_only_doc_uses_ocr_fallback(self, monkeypatch):
        monkeypatch.setattr(
            DocExtractorProxy, "_run_antiword", staticmethod(lambda p: ("[pic][pic]\n[pic]\n", False))
        )
        monkeypatch.setattr(doc_ocr, "extract_doc_images", lambda p: [_image(), _image(120, 90)])
        monkeypatch.setattr(doc_ocr, "get_ocr_engine", lambda: object())
        monkeypatch.setattr(doc_ocr, "ocr_image_bytes", lambda blob: ["识别文本行"])

        md = asyncio.run(DocExtractorProxy().extract_to_markdown("fake.doc"))
        assert "[pic]" not in md
        assert "识别文本行" in md
        assert "## 图片 1" in md

    def test_placeholder_doc_keeps_antiword_captions(self, monkeypatch):
        monkeypatch.setattr(
            DocExtractorProxy,
            "_run_antiword",
            staticmethod(lambda p: ("成功客户充值1\n[pic][pic]\n[pic]\n", False)),
        )
        monkeypatch.setattr(doc_ocr, "extract_doc_images", lambda p: [_image()])
        monkeypatch.setattr(doc_ocr, "get_ocr_engine", lambda: object())
        monkeypatch.setattr(doc_ocr, "ocr_image_bytes", lambda blob: ["OCR内容"])

        md = asyncio.run(DocExtractorProxy().extract_to_markdown("fake.doc"))
        assert md.startswith("成功客户充值1")
        assert "OCR内容" in md

    def test_antiword_failure_without_images_reports_error(self, monkeypatch):
        monkeypatch.setattr(DocExtractorProxy, "_run_antiword", staticmethod(lambda p: (None, True)))
        monkeypatch.setattr(doc_ocr, "extract_doc_images", lambda p: [])
        md = asyncio.run(DocExtractorProxy().extract_to_markdown("fake.doc"))
        assert "Error parsing DOC" in md

    def test_antiword_lenient_decoding_survives_bad_bytes(self, monkeypatch):
        """antiword output with stray non-UTF-8 bytes must not crash extraction.

        Real antiword emits UTF-8 (CP936 mapped via its tables) with the
        occasional malformed byte; strict text=True decoding used to blow up
        the whole extraction with UnicodeDecodeError.
        """
        import subprocess

        payload = "第二天性格灌输，家庭摸底".encode("utf-8") + b"\xed\xed" + "。讲述故事".encode("utf-8")

        class _Result:
            returncode = 0
            stdout = payload
            stderr = b""

        monkeypatch.setattr(subprocess, "run", lambda *a, **kw: _Result())
        text, failed = DocExtractorProxy._run_antiword("fake.doc")
        assert not failed
        assert "第二天性格灌输" in text
        assert "讲述故事" in text


# ---------------------------------------------------------------------------
# DOCX walker: tables, headers/footers, encrypted diagnostics, image fallback
# ---------------------------------------------------------------------------

def _build_docx(path, *, with_table=True, with_header=False, image_blob=None):
    import docx

    doc = docx.Document()
    doc.add_heading("培训要素", level=1)
    doc.add_paragraph("第一天的任务是建立信任。")
    if with_table:
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "要素"
        table.cell(0, 1).text = "话术"
        table.cell(1, 0).text = "家庭住址"
        table.cell(1, 1).text = "带竖线|的单元格"
    if with_header:
        doc.sections[0].header.paragraphs[0].text = "内部培训资料"
    if image_blob is not None:
        doc.add_picture(io.BytesIO(image_blob))
    doc.save(path)


class TestDocxExtractor:
    def test_tables_and_headers_are_included(self, tmp_path):
        path = str(tmp_path / "t.docx")
        _build_docx(path, with_table=True, with_header=True)

        md = asyncio.run(DocxExtractor().extract_to_markdown(path))
        assert "# 培训要素" in md
        assert "第一天的任务是建立信任" in md
        assert "| 要素 | 话术 |" in md            # table header row
        assert "| 家庭住址 |" in md
        assert "带竖线\\|的单元格" in md       # pipe escaping
        assert "【页眉】内部培训资料" in md

    def test_textbox_content_is_swept_from_xml(self, tmp_path):
        path = str(tmp_path / "tx.docx")
        _build_docx(path, with_table=False)
        # Inject a text box paragraph into the saved package.
        with zipfile.ZipFile(path) as zf:
            names = zf.namelist()
            blobs = {n: zf.read(n) for n in names}
        xml = blobs["word/document.xml"].decode("utf-8")
        injected = xml.replace(
            "</w:body>",
            "<w:p><w:r><w:pict><w:txbxContent><w:p><w:r>"
            "<w:t>文本框里的隐藏话术</w:t>"
            "</w:r></w:p></w:txbxContent></w:pict></w:r></w:p></w:body>",
        )
        blobs["word/document.xml"] = injected.encode("utf-8")
        with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zf:
            for name, blob in blobs.items():
                zf.writestr(name, blob)

        md = asyncio.run(DocxExtractor().extract_to_markdown(path))
        assert "【文本框】" in md
        assert "文本框里的隐藏话术" in md

    def test_encrypted_compound_doc_yields_diagnostic(self, tmp_path):
        path = str(tmp_path / "encrypted.docx")
        with open(path, "wb") as fh:
            fh.write(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 64)

        md = asyncio.run(DocxExtractor().extract_to_markdown(path))
        assert "加密" in md
        assert "无法提取" in md

    def test_image_only_docx_triggers_ocr_fallback(self, tmp_path, monkeypatch):
        path = str(tmp_path / "imgs.docx")
        _build_docx(path, with_table=False, image_blob=_png_blob())
        monkeypatch.setattr(doc_ocr, "get_ocr_engine", lambda: object())
        monkeypatch.setattr(doc_ocr, "ocr_image_bytes", lambda blob: ["截图聊天记录行"])

        md = asyncio.run(DocxExtractor().extract_to_markdown(path))
        assert "内嵌图片" in md
        assert "截图聊天记录行" in md
